"""ATS-safe DOCX rendering: single column, standard headings, no tables/text boxes."""
import datetime

from docx import Document

CONTACT_LINE = "Jacob V. Nelson | Richmond, VA (Remote) | jacob.v.nelson11@gmail.com"


def render_resume_docx(path: str, tailored: dict, skills: list[str], certifications: list[str], education: list[dict]) -> None:
    doc = Document()
    doc.add_heading("Jacob V. Nelson", level=0)
    doc.add_paragraph(CONTACT_LINE)

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(tailored["summary"])

    doc.add_heading("Experience", level=1)
    for role in tailored["experience"]:
        p = doc.add_paragraph()
        run = p.add_run(f"{role['title']} - {role['company']}")
        run.bold = True
        for bullet in role["bullets"]:
            doc.add_paragraph(bullet, style="List Bullet")

    doc.add_heading("Skills", level=1)
    doc.add_paragraph(", ".join(skills))

    doc.add_heading("Certifications", level=1)
    doc.add_paragraph(", ".join(certifications))

    doc.add_heading("Education", level=1)
    for edu in education:
        doc.add_paragraph(f"{edu['degree']}, {edu['school']} ({edu['gradDate']})")

    doc.save(path)


def render_cover_letter_docx(path: str, cover_letter: str, job_title: str, company: str) -> None:
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("Jacob V. Nelson")
    run.bold = True
    doc.add_paragraph(CONTACT_LINE)
    doc.add_paragraph(datetime.date.today().strftime("%B %d, %Y"))

    p = doc.add_paragraph()
    run = p.add_run(f"Re: {job_title} at {company}")
    run.bold = True

    for para in [p for p in cover_letter.split("\n") if p.strip()]:
        doc.add_paragraph(para)

    doc.save(path)
